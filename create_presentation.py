import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def add_run_text(paragraph, text):
    run = paragraph.add_run()
    run.text = text
    return run

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)      # #0F172A (Navy Dark)
    CARD_DARK = RGBColor(30, 41, 59)    # #1E293B (Card Dark)
    LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC (Clean Light)
    CARD_LIGHT = RGBColor(255, 255, 255)# White
    TEAL_BRAND = RGBColor(15, 118, 110) # #0F766E
    TEAL_LIGHT = RGBColor(240, 253, 250)# #F0FDFA
    ACCENT_CYAN = RGBColor(13, 148, 136)# #0D9488
    GOLD = RGBColor(217, 119, 6)        # #D97706
    TEXT_DARK = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139)

    blank_slide_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_card(slide, left, top, width, height, bg_color=CARD_LIGHT, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, section_tag, title_text, dark_mode=False):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        r0 = add_run_text(p0, section_tag.upper())
        r0.font.name = 'Calibri'
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = TEAL_BRAND if not dark_mode else ACCENT_CYAN

        tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf2 = tx_box2.text_frame
        tf2.word_wrap = True
        p1 = tf2.paragraphs[0]
        r1 = add_run_text(p1, title_text)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(24)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK if not dark_mode else RGBColor(248, 250, 252)

    # =========================================================================
    # SLIDE 1: COVER SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1, DARK_BG)
    add_card(s1, 0.8, 1.2, 11.733, 5.1, CARD_DARK, TEAL_BRAND)

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = add_run_text(p, "MEDPLUS PHARMACY POS")
    r.font.name = 'Calibri'
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf.add_paragraph()
    r2 = add_run_text(p2, "Next-Generation Point of Sale & FEFO Inventory Engine for Healthcare Retail")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(16)
    r2.font.color.rgb = ACCENT_CYAN

    tb_info = s1.shapes.add_textbox(Inches(1.2), Inches(3.2), Inches(11), Inches(2.5))
    tf_i = tb_info.text_frame
    
    items = [
        ("Presenter / Student:", "Ralph Osiko (Reg. No: 23/08138)"),
        ("Degree Program:", "Bachelor of Science in Information Technology"),
        ("Project Supervisor:", "Mr. Jacob Oloo"),
        ("Institution:", "School of Technology, KCA University"),
        ("Target Market:", "Community Pharmacies & Retail Healthcare Outlets in Kenya")
    ]
    for lbl, val in items:
        p_i = tf_i.add_paragraph()
        p_i.space_after = Pt(6)
        r_l = add_run_text(p_i, f"{lbl:<24} ")
        r_l.font.bold = True
        r_l.font.size = Pt(13)
        r_l.font.color.rgb = RGBColor(203, 213, 225)
        r_v = add_run_text(p_i, val)
        r_v.font.size = Pt(13)
        r_v.font.color.rgb = RGBColor(255, 255, 255)

    # =========================================================================
    # SLIDE 2: EXECUTIVE OVERVIEW & PITCH
    # =========================================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2, LIGHT_BG)
    add_header(s2, "Executive Overview", "Modernizing Retail Pharmacy Operations in Kenya")

    m_data = [
        ("Ksh 120B+", "Kenyan Retail Pharma Market", "Rapidly expanding community healthcare sector"),
        ("< 500ms", "POS Checkout Speed", "Instant checkout with automated M-Pesa STK Push"),
        ("100% FEFO", "Expiry Prevention", "Zero sale of expired drugs via batch lock controls")
    ]
    for idx, (stat, label, desc) in enumerate(m_data):
        left = 0.8 + idx * 3.95
        add_card(s2, left, 1.6, 3.8, 1.8, CARD_LIGHT, RGBColor(203, 213, 225))
        tb_m = s2.shapes.add_textbox(Inches(left + 0.2), Inches(1.7), Inches(3.4), Inches(1.5))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        
        p_s = tf_m.paragraphs[0]
        r_s = add_run_text(p_s, stat)
        r_s.font.size = Pt(28)
        r_s.font.bold = True
        r_s.font.color.rgb = TEAL_BRAND
        
        p_l = tf_m.add_paragraph()
        r_l = add_run_text(p_l, label)
        r_l.font.size = Pt(13)
        r_l.font.bold = True
        r_l.font.color.rgb = TEXT_DARK
        
        p_d = tf_m.add_paragraph()
        r_d = add_run_text(p_d, desc)
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = TEXT_MUTED

    add_card(s2, 0.8, 3.6, 11.733, 3.2, CARD_DARK)
    tb_pitch = s2.shapes.add_textbox(Inches(1.1), Inches(3.8), Inches(11.1), Inches(2.8))
    tf_p = tb_pitch.text_frame
    tf_p.word_wrap = True
    
    p_head = tf_p.paragraphs[0]
    r_h = add_run_text(p_head, "Core Value Proposition for Investors & Pharmacies:")
    r_h.font.size = Pt(16)
    r_h.font.bold = True
    r_h.font.color.rgb = ACCENT_CYAN

    pitch_bullets = [
        ("High-Performance Go Micro-Monolith:", "Compiled backend providing enterprise stability without recurring software license fees."),
        ("Native M-Pesa Daraja 2.0 Integration:", "Direct Safaricom API connection for instant mobile money payment verification at checkout."),
        ("PPB Regulatory Compliance:", "Enforces doctor prescription verification and strict First-Expired, First-Out (FEFO) stock rotation."),
        ("Empirical Reliability:", "Validated with a 100% PASS rate across 51 GUI, Database, and System Output test benchmarks.")
    ]
    for b_title, b_desc in pitch_bullets:
        p_b = tf_p.add_paragraph()
        p_b.space_before = Pt(4)
        r_bt = add_run_text(p_b, f"• {b_title} ")
        r_bt.font.bold = True
        r_bt.font.size = Pt(12)
        r_bt.font.color.rgb = RGBColor(255, 255, 255)
        r_bd = add_run_text(p_b, b_desc)
        r_bd.font.size = Pt(12)
        r_bd.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3, LIGHT_BG)
    add_header(s3, "Problem Statement", "Critical Operational Pain Points in Retail Pharmacies")

    problems = [
        ("Expired Drug Dispensing Risk", "Manual paper ledgers cannot track batch shelf-life across hundreds of medicines, risking accidental dispensing of expired drugs and PPB penalties."),
        ("Stock Shrinkage & Loss", "Lack of real-time stock deduction and automated reorder alerts leads to inventory shrinkage, theft, and revenue leakage."),
        ("Rx Non-Compliance", "Dispensing Prescription-Only Medication (POM) without verified doctor details risks patient safety and regulatory sanctions."),
        ("Slow Checkout & Cash Errors", "Manual cash calculations and manual M-Pesa status checks slow down sales processing and cause daily balancing discrepancies.")
    ]
    for idx, (p_title, p_desc) in enumerate(problems):
        row = idx // 2
        col = idx % 2
        left = 0.8 + col * 5.95
        top = 1.6 + row * 2.7
        add_card(s3, left, top, 5.7, 2.4, CARD_LIGHT, RGBColor(226, 232, 240))
        
        tb_p = s3.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(5.3), Inches(2.0))
        tf_pr = tb_p.text_frame
        tf_pr.word_wrap = True
        
        p_pt = tf_pr.paragraphs[0]
        r_pt = add_run_text(p_pt, f"0{idx+1}. {p_title}")
        r_pt.font.size = Pt(15)
        r_pt.font.bold = True
        r_pt.font.color.rgb = RGBColor(185, 28, 28)
        
        p_pd = tf_pr.add_paragraph()
        p_pd.space_before = Pt(6)
        r_pd = add_run_text(p_pd, p_desc)
        r_pd.font.size = Pt(11.5)
        r_pd.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 4: SYSTEM OBJECTIVES
    # =========================================================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4, LIGHT_BG)
    add_header(s4, "Project Scope", "Core Project Objectives & Key Deliverables")

    add_card(s4, 0.8, 1.6, 4.0, 5.2, TEAL_BRAND)
    tb_obj = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.6), Inches(4.8))
    tf_o = tb_obj.text_frame
    tf_o.word_wrap = True
    
    p_om = tf_o.paragraphs[0]
    r_om = add_run_text(p_om, "MAIN OBJECTIVE")
    r_om.font.size = Pt(14)
    r_om.font.bold = True
    r_om.font.color.rgb = ACCENT_CYAN

    p_om2 = tf_o.add_paragraph()
    p_om2.space_before = Pt(10)
    r_om2 = add_run_text(p_om2, "To design, implement, and evaluate a web-based Pharmacy POS System for MedPlus Pharmacy that automates sales checkout, enforces FEFO inventory expiry control, validates doctor prescriptions, and processes M-Pesa payments.")
    r_om2.font.size = Pt(13.5)
    r_om2.font.color.rgb = RGBColor(255, 255, 255)

    sp_objectives = [
        "1. Dynamic Auth & RBAC (Bcrypt, Admin/Pharmacist/Cashier roles)",
        "2. FEFO Inventory Control (Batch tracking & automated reorder alerts)",
        "3. Fast POS Checkout (Barcode scanning & M-Pesa STK Push API)",
        "4. Prescription Verification (Doctor & patient validation prior to sale)",
        "5. Supply Chain Mgmt (Supplier profiles & Purchase Order receiving)",
        "6. Empirical Validation (100% Pass across GUI, DB & Output benchmarks)"
    ]
    for idx, sp_text in enumerate(sp_objectives):
        row = idx // 2
        col = idx % 2
        left = 5.1 + col * 3.7
        top = 1.6 + row * 1.7
        add_card(s4, left, top, 3.5, 1.5, CARD_LIGHT, RGBColor(203, 213, 225))
        
        tb_sp = s4.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.15), Inches(3.2), Inches(1.2))
        tf_sp = tb_sp.text_frame
        tf_sp.word_wrap = True
        p_sp = tf_sp.paragraphs[0]
        r_sp = add_run_text(p_sp, sp_text)
        r_sp.font.size = Pt(11.5)
        r_sp.font.bold = True
        r_sp.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 5: COMPETITIVE POSITIONING
    # =========================================================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5, LIGHT_BG)
    add_header(s5, "Competitive Positioning", "MedPlus POS vs Existing Solutions")

    col_w = [2.6, 2.2, 2.2, 2.2, 2.5]
    table_shape = s5.shapes.add_table(5, 5, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1))
    table = table_shape.table

    for i, w in enumerate(col_w):
        table.columns[i].width = Inches(w)

    table_data = [
        ["Feature / Capability", "OpenMRS / Bahmni", "Odoo Pharmacy", "QuickBooks POS", "MedPlus POS"],
        ["Architecture & Stack", "Java Enterprise (Heavy)", "Python / PostgreSQL", "Desktop C# / SQL", "Go Micro-Monolith (Fast)"],
        ["M-Pesa STK Push API", "No (Custom plugin needed)", "Paid Third-Party Plugin", "No Native Integration", "Native Daraja 2.0 API"],
        ["FEFO Expiry Control", "Partial Batch Tracking", "Yes (Enterprise Edition)", "No (FIFO/LIFO only)", "Automated FEFO & Alerts"],
        ["Recurring Licensing Fee", "Open Source (High Setup)", "High Monthly / User Fee", "Perpetual / Legacy Fee", "Zero Recurring License Fee"]
    ]

    for row_idx, row_items in enumerate(table_data):
        for col_idx, text_val in enumerate(row_items):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p_c = cell.text_frame.paragraphs[0]
            p_c.alignment = PP_ALIGN.LEFT
            r_c = add_run_text(p_c, text_val)
            r_c.font.name = 'Calibri'
            r_c.font.size = Pt(11)
            
            if row_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = DARK_BG
                r_c.font.bold = True
                r_c.font.color.rgb = RGBColor(255, 255, 255)
            elif col_idx == 4:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TEAL_LIGHT
                r_c.font.bold = True
                r_c.font.color.rgb = TEAL_BRAND
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 == 1 else CARD_LIGHT
                r_c.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 6: TECHNICAL ARCHITECTURE
    # =========================================================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6, LIGHT_BG)
    add_header(s6, "Technical Architecture", "High-Concurrency 3-Tier Micro-Monolith Architecture")

    arch_layers = [
        ("CLIENT TIER", "Single-Page App (SPA)", "Vanilla HTML5 / CSS3 / JavaScript with Inter & Outfit fonts. Zero heavy client frameworks for lightning-fast DOM loading."),
        ("APPLICATION TIER", "Go RESTful Micro-Monolith", "Go 1.22+ backend compiled binary (`net/http` router, bcrypt auth, M-Pesa STK Push handler, FEFO rules engine). Response time < 500ms."),
        ("DATABASE TIER", "PostgreSQL 16/18 Relational DB", "11 normalized tables in 3NF (`users`, `products`, `sales`, `prescriptions`, `suppliers`, `mpesa_payments`, `role_permissions`).")
    ]
    for idx, (tier_name, subtitle, desc) in enumerate(arch_layers):
        top = 1.6 + idx * 1.8
        add_card(s6, 0.8, top, 11.733, 1.5, CARD_LIGHT, RGBColor(203, 213, 225))
        
        add_card(s6, 0.95, top + 0.15, 2.5, 1.2, TEAL_BRAND)
        tb_t = s6.shapes.add_textbox(Inches(1.0), Inches(top + 0.3), Inches(2.4), Inches(0.9))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        r_t = add_run_text(p_t, tier_name)
        r_t.font.size = Pt(14)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(255, 255, 255)
        
        tb_d = s6.shapes.add_textbox(Inches(3.6), Inches(top + 0.2), Inches(8.7), Inches(1.1))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_sub = tf_d.paragraphs[0]
        r_sub = add_run_text(p_sub, subtitle)
        r_sub.font.size = Pt(13.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = TEXT_DARK
        
        p_ds = tf_d.add_paragraph()
        p_ds.space_before = Pt(4)
        r_ds = add_run_text(p_ds, desc)
        r_ds.font.size = Pt(11)
        r_ds.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: KEY MODULES
    # =========================================================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7, LIGHT_BG)
    add_header(s7, "Feature Highlights", "End-to-End Pharmacy Operational Modules")

    modules = [
        ("Authentication & RBAC", "Bcrypt password protection & dynamic permission matrix across Admin, Pharmacist, Cashier roles."),
        ("FEFO Inventory Control", "Real-time stock deduction, batch expiry tracking, and automated reorder warning alerts."),
        ("POS Checkout & M-Pesa", "Barcode cart management, cash change calculation, and instant M-Pesa Daraja 2.0 STK Push."),
        ("Prescription Dispensing", "Doctor Rx validation workflow preventing unverified dispensing of prescription-only drugs."),
        ("Supply Chain & POs", "Supplier management and Purchase Order receiving into inventory catalog."),
        ("Security Audit & DB Token", "Timestamped security logs and 10-minute access token (`X-Database-Access-Token`) protection.")
    ]
    for idx, (m_title, m_desc) in enumerate(modules):
        row = idx // 3
        col = idx % 3
        left = 0.8 + col * 3.95
        top = 1.6 + row * 2.7
        add_card(s7, left, top, 3.8, 2.4, CARD_LIGHT, RGBColor(226, 232, 240))
        
        tb_mod = s7.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(3.4), Inches(2.0))
        tf_m = tb_mod.text_frame
        tf_m.word_wrap = True
        
        p_mt = tf_m.paragraphs[0]
        r_mt = add_run_text(p_mt, m_title)
        r_mt.font.size = Pt(14.5)
        r_mt.font.bold = True
        r_mt.font.color.rgb = TEAL_BRAND
        
        p_md = tf_m.add_paragraph()
        p_md.space_before = Pt(6)
        r_md = add_run_text(p_md, m_desc)
        r_md.font.size = Pt(11)
        r_md.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 8: TESTING & VERIFICATION
    # =========================================================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8, LIGHT_BG)
    add_header(s8, "Testing & Verification", "100% Pass Rate Across All Validation Tiers")

    test_tiers = [
        ("24 GUI Test Cases", "100% PASS Rate", "Validated login redirects, cart updates, Rx modal prompts, M-Pesa spinner, product catalog saves, and RBAC toggles."),
        ("15 DB Test Cases", "100% PASS Rate", "Verified UNIQUE barcode constraints, foreign key cascades, atomic checkout transactions, and DB security token checks."),
        ("12 Output Test Cases", "100% PASS Rate", "Verified 16% VAT receipt calculations, M-Pesa STK callbacks, daily sales summaries, and timestamped audit logs.")
    ]
    for idx, (tier_t, pass_lbl, tier_desc) in enumerate(test_tiers):
        left = 0.8 + idx * 3.95
        add_card(s8, left, 1.6, 3.8, 5.2, CARD_LIGHT, RGBColor(203, 213, 225))
        
        add_card(s8, left + 0.2, 1.8, 3.4, 1.1, TEAL_BRAND)
        tb_th = s8.shapes.add_textbox(Inches(left + 0.2), Inches(1.9), Inches(3.4), Inches(0.9))
        tf_th = tb_th.text_frame
        tf_th.word_wrap = True
        p_th = tf_th.paragraphs[0]
        p_th.alignment = PP_ALIGN.CENTER
        r_th = add_run_text(p_th, tier_t)
        r_th.font.size = Pt(16)
        r_th.font.bold = True
        r_th.font.color.rgb = RGBColor(255, 255, 255)
        
        p_pl = tf_th.add_paragraph()
        p_pl.alignment = PP_ALIGN.CENTER
        r_pl = add_run_text(p_pl, pass_lbl)
        r_pl.font.size = Pt(13)
        r_pl.font.bold = True
        r_pl.font.color.rgb = GOLD

        tb_td = s8.shapes.add_textbox(Inches(left + 0.3), Inches(3.1), Inches(3.2), Inches(3.5))
        tf_td = tb_td.text_frame
        tf_td.word_wrap = True
        p_td = tf_td.paragraphs[0]
        r_td = add_run_text(p_td, tier_desc)
        r_td.font.size = Pt(11.5)
        r_td.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 9: INVESTOR ROI & VALUE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9, DARK_BG)
    add_header(s9, "Investor ROI & Value", "Compelling Business Model & Market Growth", dark_mode=True)

    roi_items = [
        ("Zero Vendor License Barriers", "Replaces expensive $1,500/year enterprise ERP subscriptions with an affordable single-deployment web binary."),
        ("100% Stock Leakage Protection", "Real-time inventory deduction and FEFO controls eliminate unaccounted stock loss and expired product disposal costs."),
        ("Fast Queue Throughput", "Transaction speed reduced to under 30 seconds per customer, increasing daily retail sales capacity by 40%."),
        ("Scalable SaaS Roadmap", "Ready for multi-branch cloud deployment, recurring maintenance contracts, and payment transaction revenue-share models.")
    ]
    for idx, (roi_t, roi_d) in enumerate(roi_items):
        row = idx // 2
        col = idx % 2
        left = 0.8 + col * 5.95
        top = 1.6 + row * 2.7
        add_card(s9, left, top, 5.7, 2.4, CARD_DARK, ACCENT_CYAN)
        
        tb_roi = s9.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(5.3), Inches(2.0))
        tf_r = tb_roi.text_frame
        tf_r.word_wrap = True
        
        p_rt = tf_r.paragraphs[0]
        r_rt = add_run_text(p_rt, roi_t)
        r_rt.font.size = Pt(15)
        r_rt.font.bold = True
        r_rt.font.color.rgb = RGBColor(255, 255, 255)
        
        p_rd = tf_r.add_paragraph()
        p_rd.space_before = Pt(6)
        r_rd = add_run_text(p_rd, roi_d)
        r_rd.font.size = Pt(11.5)
        r_rd.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 10: SUMMARY & FUTURE ROADMAP
    # =========================================================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10, LIGHT_BG)
    add_header(s10, "Summary & Roadmap", "Project Conclusions & Next-Phase Enhancements")

    add_card(s10, 0.8, 1.6, 5.7, 5.2, CARD_LIGHT, RGBColor(203, 213, 225))
    tb_c = s10.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    
    p_ct = tf_c.paragraphs[0]
    r_ct = add_run_text(p_ct, "Project Conclusions:")
    r_ct.font.size = Pt(16)
    r_ct.font.bold = True
    r_ct.font.color.rgb = TEAL_BRAND

    conclusions = [
        "Successfully developed and tested a Go/PostgreSQL Pharmacy POS system for MedPlus Pharmacy.",
        "Automated checkout, prescription validation, and M-Pesa STK Push payment processing.",
        "Eliminated expired medicine sales through strict FEFO batch controls.",
        "Achieved 100% empirical pass rate across 51 comprehensive test benchmarks."
    ]
    for c_text in conclusions:
        p_cd = tf_c.add_paragraph()
        p_cd.space_before = Pt(8)
        r_cd = add_run_text(p_cd, f"✔  {c_text}")
        r_cd.font.size = Pt(11.5)
        r_cd.font.color.rgb = TEXT_DARK

    add_card(s10, 6.8, 1.6, 5.7, 5.2, CARD_LIGHT, RGBColor(203, 213, 225))
    tb_rm = s10.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_rm = tb_rm.text_frame
    tf_rm.word_wrap = True
    
    p_rmt = tf_rm.paragraphs[0]
    r_rmt = add_run_text(p_rmt, "Future Product Roadmap:")
    r_rmt.font.size = Pt(16)
    r_rmt.font.bold = True
    r_rmt.font.color.rgb = ACCENT_CYAN

    roadmap = [
        ("AI-Driven Demand Forecasting:", "Incorporate predictive models to forecast seasonal antibiotic and medication demand spikes."),
        ("Kenya SHA / NHIF Claims API:", "Direct digital insurance claim pre-authorization at checkout."),
        ("Multi-Branch Cloud PWA Sync:", "Offline-first Progressive Web App sync connecting remote pharmacy branches.")
    ]
    for rm_title, rm_desc in roadmap:
        p_rmd = tf_rm.add_paragraph()
        p_rmd.space_before = Pt(10)
        r_rmt2 = add_run_text(p_rmd, f"🚀 {rm_title} ")
        r_rmt2.font.bold = True
        r_rmt2.font.size = Pt(11.5)
        r_rmt2.font.color.rgb = TEXT_DARK
        r_rmd2 = add_run_text(p_rmd, rm_desc)
        r_rmd2.font.size = Pt(11)
        r_rmd2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 11: THANK YOU
    # =========================================================================
    s11 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s11, DARK_BG)

    add_card(s11, 1.5, 1.5, 10.333, 4.5, CARD_DARK, TEAL_BRAND)

    tb_end = s11.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.7), Inches(3.9))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True
    
    p_e1 = tf_end.paragraphs[0]
    p_e1.alignment = PP_ALIGN.CENTER
    r_e1 = add_run_text(p_e1, "THANK YOU!")
    r_e1.font.size = Pt(40)
    r_e1.font.bold = True
    r_e1.font.color.rgb = RGBColor(255, 255, 255)

    p_e2 = tf_end.add_paragraph()
    p_e2.alignment = PP_ALIGN.CENTER
    p_e2.space_before = Pt(8)
    r_e2 = add_run_text(p_e2, "Questions & Discussion")
    r_e2.font.size = Pt(20)
    r_e2.font.bold = True
    r_e2.font.color.rgb = ACCENT_CYAN

    p_e3 = tf_end.add_paragraph()
    p_e3.alignment = PP_ALIGN.CENTER
    p_e3.space_before = Pt(24)
    r_e3 = add_run_text(p_e3, "Ralph Osiko (Student, 23/08138)  |  Supervisor: Mr. Jacob Oloo\nSchool of Technology, KCA University")
    r_e3.font.size = Pt(14)
    r_e3.font.color.rgb = RGBColor(203, 213, 225)

    current_dir = os.path.dirname(os.path.abspath(__file__))
    pptx_path = os.path.join(current_dir, "Final_Presentation.pptx")
    prs.save(pptx_path)
    print(f"Successfully generated PowerPoint presentation at: {pptx_path}")

if __name__ == '__main__':
    create_deck()
