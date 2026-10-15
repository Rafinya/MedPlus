import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def add_run_text(paragraph, text):
    run = paragraph.add_run()
    run.text = text
    return run

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)        # #0F172A (Navy Dark)
    CARD_DARK = RGBColor(30, 41, 59)      # #1E293B (Card Dark)
    LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC (Clean Light)
    CARD_LIGHT = RGBColor(255, 255, 255)  # White
    TEAL_BRAND = RGBColor(15, 118, 110)   # #0F766E
    TEAL_LIGHT = RGBColor(240, 253, 250)  # #F0FDFA
    ACCENT_CYAN = RGBColor(13, 148, 136)  # #0D9488
    GOLD = RGBColor(217, 119, 6)          # #D97706
    RED_ALERT = RGBColor(185, 28, 28)     # #B91C1C
    TEXT_DARK = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139)
    TEXT_WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6]

    def set_bg(slide, color):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_card(slide, left, top, width, height, bg_color=CARD_LIGHT, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, section_tag, title_text, dark_mode=False):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        r0 = add_run_text(p0, section_tag.upper())
        r0.font.name = 'Calibri'
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = TEAL_BRAND if not dark_mode else ACCENT_CYAN

        tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf2 = tx_box2.text_frame
        tf2.word_wrap = True
        p1 = tf2.paragraphs[0]
        r1 = add_run_text(p1, title_text)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(24)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_DARK if not dark_mode else RGBColor(248, 250, 252)

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, DARK_BG)
    add_card(s1, 0.8, 1.2, 11.733, 5.1, CARD_DARK, TEAL_BRAND)

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11), Inches(1.2))
    tf1 = tb1.text_frame
    p1 = tf1.paragraphs[0]
    r1 = add_run_text(p1, "MEDPLUS PHARMACY POINT OF SALE (POS)\nMANAGEMENT SYSTEM")
    r1.font.name = 'Calibri'
    r1.font.size = Pt(30)
    r1.font.bold = True
    r1.font.color.rgb = TEXT_WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.space_before = Pt(8)
    r1_sub = add_run_text(p1_sub, "Final Year Project II - Final Examination Presentation")
    r1_sub.font.name = 'Calibri'
    r1_sub.font.size = Pt(16)
    r1_sub.font.color.rgb = ACCENT_CYAN

    tb_info = s1.shapes.add_textbox(Inches(1.2), Inches(3.4), Inches(11), Inches(2.5))
    tf_info = tb_info.text_frame
    meta_items = [
        ("Presenter:", "Ralph Osiko (Reg. No: 23/08138)"),
        ("Course:", "Bachelor of Science in Information Technology"),
        ("Supervisor:", "Mr. Jacob Oloo"),
        ("Institution:", "School of Technology, KCA University"),
        ("Submission Date:", "July 2026")
    ]
    for lbl, val in meta_items:
        p_m = tf_info.add_paragraph()
        p_m.space_after = Pt(4)
        r_l = add_run_text(p_m, f"{lbl:<18} ")
        r_l.font.bold = True
        r_l.font.size = Pt(13)
        r_l.font.color.rgb = RGBColor(203, 213, 225)
        r_v = add_run_text(p_m, val)
        r_v.font.size = Pt(13)
        r_v.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 2: OUTLINE (EXACT SLIDE MATCHING USER TEMPLATE)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, DARK_BG)
    add_header(s2, "Presentation Structure", "Outline of Topics Covered", dark_mode=True)

    add_card(s2, 0.8, 1.6, 11.733, 5.2, CARD_DARK, ACCENT_CYAN)

    tb_out = s2.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.8))
    tf_out = tb_out.text_frame
    tf_out.word_wrap = True

    outline_items = [
        ("•  Introduction", "Overview of retail pharmacy operations & digital POS transformation"),
        ("•  Background", "Operational risks, FEFO batch tracking, and PPB healthcare regulations"),
        ("•  Problem statement", "Expired drug dispensing, inventory theft, Rx non-compliance, slow queues"),
        ("•  System objectives", "Main objective & specific objectives 1 through 6"),
        ("•  Significance of Study", "Patient health safety, financial accountability, queue performance"),
        ("•  Literature review", "Comparative evaluation of existing POS systems & technical research gaps"),
        ("•  Methodology", "Data collection, Waterfall model, 3-tier Go/PostgreSQL architecture"),
        ("•  Results (screenshots of the developed systems and explanations)", "Live UI demonstration, POS checkout cart, M-Pesa modal & achieved objectives")
    ]

    for title, desc in outline_items:
        p_o = tf_out.add_paragraph()
        p_o.space_before = Pt(5)
        r_ot = add_run_text(p_o, f"{title:<40} ")
        r_ot.font.name = 'Calibri'
        r_ot.font.size = Pt(13.5)
        r_ot.font.bold = True
        r_ot.font.color.rgb = TEXT_WHITE
        
        r_od = add_run_text(p_o, f"—  {desc}")
        r_od.font.name = 'Calibri'
        r_od.font.size = Pt(12)
        r_od.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 3: INTRODUCTION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, LIGHT_BG)
    add_header(s3, "Section 01", "Introduction to Pharmacy Operations & Digital POS Transformation")

    intro_cards = [
        ("Primary Healthcare Touchpoints", "Community retail pharmacies in Kenya dispense life-saving prescription medicines, Over-The-Counter (OTC) drugs, and wellness products to thousands of patients daily."),
        ("Expanding Operational Volume", "Rapid growth in urban healthcare retail has significantly increased transaction frequency, drug batch processing, and daily customer consultation volume."),
        ("Modern Web Architecture Shift", "Transitioning from legacy manual ledgers to an integrated web-based Point of Sale (POS) system powered by Go, PostgreSQL, and Vanilla JavaScript.")
    ]
    for idx, (title, desc) in enumerate(intro_cards):
        left = 0.8 + idx * 3.95
        add_card(s3, left, 1.6, 3.8, 5.2, CARD_LIGHT, RGBColor(203, 213, 225))
        
        add_card(s3, left + 0.2, 1.8, 3.4, 1.1, TEAL_BRAND)
        tb_ic = s3.shapes.add_textbox(Inches(left + 0.2), Inches(1.9), Inches(3.4), Inches(0.9))
        tf_ic = tb_ic.text_frame
        tf_ic.word_wrap = True
        p_ic = tf_ic.paragraphs[0]
        p_ic.alignment = PP_ALIGN.CENTER
        r_ic = add_run_text(p_ic, title)
        r_ic.font.size = Pt(14)
        r_ic.font.bold = True
        r_ic.font.color.rgb = TEXT_WHITE

        tb_id = s3.shapes.add_textbox(Inches(left + 0.3), Inches(3.1), Inches(3.2), Inches(3.5))
        tf_id = tb_id.text_frame
        tf_id.word_wrap = True
        p_id = tf_id.paragraphs[0]
        r_id = add_run_text(p_id, desc)
        r_id.font.size = Pt(12)
        r_id.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 4: BACKGROUND
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, LIGHT_BG)
    add_header(s4, "Section 02", "Background & Pharmaceutical Regulatory Environment")

    bg_cards = [
        ("Pharmaceutical Inventory Nuances", "Unlike standard retail goods, pharmaceuticals carry strict shelf-life expiration dates, batch lot numbers, storage mandates, and regulatory controls under the Pharmacy and Poisons Act (Cap 244)."),
        ("First-Expired, First-Out (FEFO)", "Pharmacies require automated FEFO stock rotation to ensure batches nearing expiry are dispensed first, avoiding inventory write-offs and patient safety hazards."),
        ("Localized Mobile Payment Ecosystem", "Integration with Safaricom M-Pesa Daraja 2.0 STK Push API provides seamless instant payment confirmation matching Kenyan retail payment habits.")
    ]
    for idx, (b_title, b_desc) in enumerate(bg_cards):
        top = 1.6 + idx * 1.8
        add_card(s4, 0.8, top, 11.733, 1.5, CARD_LIGHT, RGBColor(203, 213, 225))
        
        add_card(s4, 0.95, top + 0.15, 2.8, 1.2, CARD_DARK)
        tb_bt = s4.shapes.add_textbox(Inches(1.0), Inches(top + 0.3), Inches(2.7), Inches(0.9))
        tf_bt = tb_bt.text_frame
        tf_bt.word_wrap = True
        p_bt = tf_bt.paragraphs[0]
        r_bt = add_run_text(p_bt, b_title)
        r_bt.font.size = Pt(13)
        r_bt.font.bold = True
        r_bt.font.color.rgb = ACCENT_CYAN
        
        tb_bd = s4.shapes.add_textbox(Inches(3.9), Inches(top + 0.2), Inches(8.4), Inches(1.1))
        tf_bd = tb_bd.text_frame
        tf_bd.word_wrap = True
        p_bd = tf_bd.paragraphs[0]
        r_bd = add_run_text(p_bd, b_desc)
        r_bd.font.size = Pt(12)
        r_bd.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 5: PROBLEM STATEMENT
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, LIGHT_BG)
    add_header(s5, "Section 03", "Problem Statement & Operational Vulnerabilities")

    problems = [
        ("01. Expired Drug Dispensing Risk", "Manual paper ledgers cannot track batch shelf-life across hundreds of medicines, risking accidental dispensing of expired drugs and PPB penalties."),
        ("02. Stock Shrinkage & Loss", "Lack of real-time stock deduction and automated reorder alerts leads to inventory shrinkage, theft, and revenue leakage."),
        ("03. Rx Prescription Non-Compliance", "Dispensing Prescription-Only Medication (POM) without verified doctor details risks patient safety and regulatory sanctions."),
        ("04. Slow Queues & Cash Errors", "Manual cash calculations and manual M-Pesa status checks slow down sales processing and cause daily balancing discrepancies.")
    ]
    for idx, (p_title, p_desc) in enumerate(problems):
        row = idx // 2
        col = idx % 2
        left = 0.8 + col * 5.95
        top = 1.6 + row * 2.7
        add_card(s5, left, top, 5.7, 2.4, CARD_LIGHT, RGBColor(226, 232, 240))
        
        tb_p = s5.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(5.3), Inches(2.0))
        tf_pr = tb_p.text_frame
        tf_pr.word_wrap = True
        
        p_pt = tf_pr.paragraphs[0]
        r_pt = add_run_text(p_pt, p_title)
        r_pt.font.size = Pt(15)
        r_pt.font.bold = True
        r_pt.font.color.rgb = RED_ALERT
        
        p_pd = tf_pr.add_paragraph()
        p_pd.space_before = Pt(6)
        r_pd = add_run_text(p_pd, p_desc)
        r_pd.font.size = Pt(11.5)
        r_pd.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 6: SYSTEM OBJECTIVES
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, LIGHT_BG)
    add_header(s6, "Section 04", "System Objectives (Main & Specific Objectives 1-6)")

    add_card(s6, 0.8, 1.6, 4.0, 5.2, TEAL_BRAND)
    tb_obj = s6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.6), Inches(4.8))
    tf_o = tb_obj.text_frame
    tf_o.word_wrap = True
    
    p_om = tf_o.paragraphs[0]
    r_om = add_run_text(p_om, "MAIN OBJECTIVE")
    r_om.font.size = Pt(14)
    r_om.font.bold = True
    r_om.font.color.rgb = ACCENT_CYAN

    p_om2 = tf_o.add_paragraph()
    p_om2.space_before = Pt(10)
    r_om2 = add_run_text(p_om2, "To design, implement, and evaluate a web-based Pharmacy POS System for MedPlus Pharmacy that automates sales checkout, enforces FEFO inventory expiry control, validates doctor prescriptions, and processes M-Pesa payments.")
    r_om2.font.size = Pt(13.5)
    r_om2.font.color.rgb = TEXT_WHITE

    sp_objectives = [
        "1. Dynamic Auth & RBAC (Bcrypt, Admin/Pharmacist/Cashier roles)",
        "2. FEFO Inventory Control (Batch tracking & automated reorder alerts)",
        "3. Fast POS Checkout (Barcode scanning & M-Pesa STK Push API)",
        "4. Prescription Verification (Doctor & patient validation prior to sale)",
        "5. Supply Chain Mgmt (Supplier profiles & Purchase Order receiving)",
        "6. Empirical Validation (100% Pass across GUI, DB & Output benchmarks)"
    ]
    for idx, sp_text in enumerate(sp_objectives):
        row = idx // 2
        col = idx % 2
        left = 5.1 + col * 3.7
        top = 1.6 + row * 1.7
        add_card(s6, left, top, 3.5, 1.5, CARD_LIGHT, RGBColor(203, 213, 225))
        
        tb_sp = s6.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.15), Inches(3.2), Inches(1.2))
        tf_sp = tb_sp.text_frame
        tf_sp.word_wrap = True
        p_sp = tf_sp.paragraphs[0]
        r_sp = add_run_text(p_sp, sp_text)
        r_sp.font.size = Pt(11.5)
        r_sp.font.bold = True
        r_sp.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 7: SIGNIFICANCE OF STUDY
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, LIGHT_BG)
    add_header(s7, "Section 05", "Significance of Study & Stakeholder Value")

    sig_cards = [
        ("Patient Health & Safety", "Prevents accidental sales of expired medication through automated FEFO blocking, safeguarding community health."),
        ("Retail Pharmacy Profitability", "Eliminates untracked inventory shrinkage, theft, and expired product write-offs, maximizing operational margins."),
        ("Fast Customer Queue Throughput", "Reduces transaction checkout times to under 30 seconds, accelerating queue flow during peak pharmacy hours."),
        ("Academic & Industry Value", "Demonstrates the high-concurrency performance of Go micro-monolith web systems for African healthcare retail.")
    ]
    for idx, (s_title, s_desc) in enumerate(sig_cards):
        row = idx // 2
        col = idx % 2
        left = 0.8 + col * 5.95
        top = 1.6 + row * 2.7
        add_card(s7, left, top, 5.7, 2.4, CARD_LIGHT, RGBColor(203, 213, 225))
        
        tb_sig = s7.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(5.3), Inches(2.0))
        tf_s = tb_sig.text_frame
        tf_s.word_wrap = True
        
        p_st = tf_s.paragraphs[0]
        r_st = add_run_text(p_st, s_title)
        r_st.font.size = Pt(16)
        r_st.font.bold = True
        r_st.font.color.rgb = TEAL_BRAND
        
        p_sd = tf_s.add_paragraph()
        p_sd.space_before = Pt(6)
        r_sd = add_run_text(p_sd, s_desc)
        r_sd.font.size = Pt(11.5)
        r_sd.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 8: LITERATURE REVIEW
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, LIGHT_BG)
    add_header(s8, "Section 06", "Literature Review & Comparative Analysis")

    col_w = [2.6, 2.2, 2.2, 2.2, 2.5]
    table_shape = s8.shapes.add_table(5, 5, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1))
    table = table_shape.table

    for i, w in enumerate(col_w):
        table.columns[i].width = Inches(w)

    table_data = [
        ["Feature / Capability", "OpenMRS / Bahmni", "Odoo Pharmacy", "QuickBooks POS", "MedPlus POS"],
        ["Architecture & Stack", "Java Enterprise (Heavy)", "Python / PostgreSQL", "Desktop C# / SQL", "Go Micro-Monolith (Fast)"],
        ["M-Pesa STK Push API", "No (Custom plugin needed)", "Paid Third-Party Plugin", "No Native Integration", "Native Daraja 2.0 API"],
        ["FEFO Expiry Control", "Partial Batch Tracking", "Yes (Enterprise Edition)", "No (FIFO/LIFO only)", "Automated FEFO & Alerts"],
        ["Recurring Licensing Fee", "Open Source (High Setup)", "High Monthly / User Fee", "Perpetual / Legacy Fee", "Zero Recurring License Fee"]
    ]

    for row_idx, row_items in enumerate(table_data):
        for col_idx, text_val in enumerate(row_items):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p_c = cell.text_frame.paragraphs[0]
            p_c.alignment = PP_ALIGN.LEFT
            r_c = add_run_text(p_c, text_val)
            r_c.font.name = 'Calibri'
            r_c.font.size = Pt(11)
            
            if row_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = DARK_BG
                r_c.font.bold = True
                r_c.font.color.rgb = TEXT_WHITE
            elif col_idx == 4:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TEAL_LIGHT
                r_c.font.bold = True
                r_c.font.color.rgb = TEAL_BRAND
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 == 1 else CARD_LIGHT
                r_c.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 9: METHODOLOGY
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9, LIGHT_BG)
    add_header(s9, "Section 07", "Methodology, Waterfall Model & System Architecture")

    meth_cards = [
        ("Data Collection", "Interviews with 4 community pharmacists, direct observation of 50 transactions, and document analysis of prescription ledgers."),
        ("Waterfall Development", "Sequential phases: (1) Requirements Analysis, (2) System Design, (3) Implementation, (4) Testing, (5) Deployment, (6) Maintenance."),
        ("3-Tier Architecture", "Client SPA (Vanilla JS), Application Tier (Go RESTful micro-monolith), Database Tier (11 normalized PostgreSQL tables in 3NF).")
    ]
    for idx, (m_title, m_desc) in enumerate(meth_cards):
        left = 0.8 + idx * 3.95
        add_card(s9, left, 1.6, 3.8, 5.2, CARD_LIGHT, RGBColor(203, 213, 225))
        
        add_card(s9, left + 0.2, 1.8, 3.4, 1.1, TEAL_BRAND)
        tb_mc = s9.shapes.add_textbox(Inches(left + 0.2), Inches(1.9), Inches(3.4), Inches(0.9))
        tf_mc = tb_mc.text_frame
        tf_mc.word_wrap = True
        p_mc = tf_mc.paragraphs[0]
        p_mc.alignment = PP_ALIGN.CENTER
        r_mc = add_run_text(p_mc, m_title)
        r_mc.font.size = Pt(14)
        r_mc.font.bold = True
        r_mc.font.color.rgb = TEXT_WHITE

        tb_md = s9.shapes.add_textbox(Inches(left + 0.3), Inches(3.1), Inches(3.2), Inches(3.5))
        tf_md = tb_md.text_frame
        tf_md.word_wrap = True
        p_md = tf_md.paragraphs[0]
        r_md = add_run_text(p_md, m_desc)
        r_md.font.size = Pt(12)
        r_md.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 10: RESULTS (SCREENSHOTS & EXPLANATIONS)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10, LIGHT_BG)
    add_header(s10, "Section 08 - Results", "Results: Screenshots of Developed System & Explanations")

    # Left Image Container
    current_dir = os.path.dirname(os.path.abspath(__file__))
    img_path = os.path.join(current_dir, "pos_screen.png")
    if os.path.exists(img_path):
        s10.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), width=Inches(6.8))

    # Right Explanation Card
    add_card(s10, 7.8, 1.6, 4.733, 5.2, CARD_LIGHT, RGBColor(203, 213, 225))
    tb_exp = s10.shapes.add_textbox(Inches(8.0), Inches(1.8), Inches(4.333), Inches(4.8))
    tf_exp = tb_exp.text_frame
    tf_exp.word_wrap = True

    p_et = tf_exp.paragraphs[0]
    r_et = add_run_text(p_et, "System UI & Features:")
    r_et.font.size = Pt(16)
    r_et.font.bold = True
    r_et.font.color.rgb = TEAL_BRAND

    features_list = [
        ("Responsive POS Cart Grid:", "Real-time drug search, barcode scanning, and fast cart quantity management."),
        ("Prescription Verification Modal:", "Prompts for Doctor Rx ID before dispensing prescription-only drugs."),
        ("M-Pesa STK Push Integration:", "Displays loading status spinner while customer inputs M-Pesa PIN."),
        ("FEFO Stock Expiry Badges:", "Visual red alerts for drugs expiring in under 30 days."),
        ("Database Admin Overview:", "Requires 10-min access token (`X-Database-Access-Token`).")
    ]
    for f_title, f_desc in features_list:
        p_fl = tf_exp.add_paragraph()
        p_fl.space_before = Pt(6)
        r_ft = add_run_text(p_fl, f"• {f_title} ")
        r_ft.font.bold = True
        r_ft.font.size = Pt(11.5)
        r_ft.font.color.rgb = TEXT_DARK
        r_fd = add_run_text(p_fl, f_desc)
        r_fd.font.size = Pt(10.5)
        r_fd.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 11: ACHIEVED OBJECTIVES & TEST METRICS
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11, LIGHT_BG)
    add_header(s11, "Section 08 - Results (Cont.)", "Achieved Objectives & Empirical Testing Results")

    test_tiers = [
        ("24 GUI Test Cases", "100% PASS Rate", "Validated login redirects, cart updates, Rx modal prompts, M-Pesa spinner, product catalog saves, and RBAC toggles."),
        ("15 DB Test Cases", "100% PASS Rate", "Verified UNIQUE barcode constraints, foreign key cascades, atomic checkout transactions, and DB security token checks."),
        ("12 Output Test Cases", "100% PASS Rate", "Verified 16% VAT receipt calculations, M-Pesa STK callbacks, daily sales summaries, and timestamped audit logs.")
    ]
    for idx, (tier_t, pass_lbl, tier_desc) in enumerate(test_tiers):
        left = 0.8 + idx * 3.95
        add_card(s11, left, 1.6, 3.8, 5.2, CARD_LIGHT, RGBColor(203, 213, 225))
        
        add_card(s11, left + 0.2, 1.8, 3.4, 1.1, TEAL_BRAND)
        tb_th = s11.shapes.add_textbox(Inches(left + 0.2), Inches(1.9), Inches(3.4), Inches(0.9))
        tf_th = tb_th.text_frame
        tf_th.word_wrap = True
        p_th = tf_th.paragraphs[0]
        p_th.alignment = PP_ALIGN.CENTER
        r_th = add_run_text(p_th, tier_t)
        r_th.font.size = Pt(16)
        r_th.font.bold = True
        r_th.font.color.rgb = TEXT_WHITE
        
        p_pl = tf_th.add_paragraph()
        p_pl.alignment = PP_ALIGN.CENTER
        r_pl = add_run_text(p_pl, pass_lbl)
        r_pl.font.size = Pt(13)
        r_pl.font.bold = True
        r_pl.font.color.rgb = GOLD

        tb_td = s11.shapes.add_textbox(Inches(left + 0.3), Inches(3.1), Inches(3.2), Inches(3.5))
        tf_td = tb_td.text_frame
        tf_td.word_wrap = True
        p_td = tf_td.paragraphs[0]
        r_td = add_run_text(p_td, tier_desc)
        r_td.font.size = Pt(11.5)
        r_td.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 12: CONCLUSION, RECOMMENDATIONS & THANK YOU
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12, DARK_BG)

    add_card(s12, 1.5, 1.2, 10.333, 5.1, CARD_DARK, TEAL_BRAND)

    tb_end = s12.shapes.add_textbox(Inches(1.8), Inches(1.4), Inches(9.7), Inches(4.5))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True
    
    p_e1 = tf_end.paragraphs[0]
    p_e1.alignment = PP_ALIGN.CENTER
    r_e1 = add_run_text(p_e1, "THANK YOU!")
    r_e1.font.size = Pt(38)
    r_e1.font.bold = True
    r_e1.font.color.rgb = TEXT_WHITE

    p_e2 = tf_end.add_paragraph()
    p_e2.alignment = PP_ALIGN.CENTER
    p_e2.space_before = Pt(6)
    r_e2 = add_run_text(p_e2, "Questions & Discussion")
    r_e2.font.size = Pt(20)
    r_e2.font.bold = True
    r_e2.font.color.rgb = ACCENT_CYAN

    p_e3 = tf_end.add_paragraph()
    p_e3.space_before = Pt(18)
    r_e3 = add_run_text(p_e3, "Key Achievements & Summary:")
    r_e3.font.size = Pt(14)
    r_e3.font.bold = True
    r_e3.font.color.rgb = GOLD

    conc_bullets = [
        "Fully operational Go/PostgreSQL Pharmacy POS system with native M-Pesa STK Push integration.",
        "Zero expired medicine sales via strict FEFO batch controls & 100% empirical test pass rate.",
        "Future Scope: AI demand forecasting, Kenya SHA/NHIF claim integration, multi-branch cloud PWA."
    ]
    for cb in conc_bullets:
        p_cb = tf_end.add_paragraph()
        p_cb.space_before = Pt(4)
        r_cb = add_run_text(p_cb, f"✔  {cb}")
        r_cb.font.size = Pt(12)
        r_cb.font.color.rgb = RGBColor(203, 213, 225)

    p_e4 = tf_end.add_paragraph()
    p_e4.alignment = PP_ALIGN.CENTER
    p_e4.space_before = Pt(20)
    r_e4 = add_run_text(p_e4, "Ralph Osiko (Student, 23/08138)  |  Supervisor: Mr. Jacob Oloo\nSchool of Technology, KCA University")
    r_e4.font.size = Pt(13)
    r_e4.font.color.rgb = RGBColor(148, 163, 184)

    # Save to both paths!
    p1 = os.path.join(current_dir, "Final_Presentation.pptx")
    p2 = os.path.join(current_dir, "Final Presentation_Template.pptx")
    prs.save(p1)
    prs.save(p2)
    print(f"Successfully generated PowerPoint presentations at:\n 1. {p1}\n 2. {p2}")

if __name__ == '__main__':
    create_presentation()
